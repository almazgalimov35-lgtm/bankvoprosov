from kivy.app import App
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.button import Button
from kivy.uix.label import Label
from kivy.uix.textinput import TextInput
from kivy.uix.scrollview import ScrollView
from kivy.uix.gridlayout import GridLayout
from kivy.metrics import dp, sp
from kivy.core.window import Window
from kivy.clock import Clock
from openpyxl import load_workbook
from collections import defaultdict

Window.clearcolor = (1, 1, 1, 1)
FILE = 'test.xlsx'
SHEET = 'Внешняя оценка знаний'
PAGE_SIZE = 10
MAX_SUGGESTIONS = 8
QUICK_ROW_1 = ['ПДК','СИЗОД','взрыв','отравление']
QUICK_ROW_2 = ['50','первая помощь','наряд','эвакуация']
QUICK_ROW_3 = ['коррозия','растворимость','контроль','замер']

wb = load_workbook(FILE, data_only=True, read_only=True)
ws = wb[SHEET] if SHEET in wb.sheetnames else wb.active
data = defaultdict(list)
all_items = []
word_counter = defaultdict(int)
for row in ws.iter_rows(min_row=2, values_only=True):
    if not row: continue
    q = str(row[1]).strip() if len(row)>1 and row[1] else ''
    if not q: continue
    code_ = str(row[0]).strip() if len(row)>0 and row[0] else ''
    a = str(row[2]).strip() if len(row)>2 and row[2] else ''
    sec = str(row[4]).strip() if len(row)>4 and row[4] else 'No section'
    data[sec].append((code_, q, a))
    all_items.append((sec, code_, q, a))
    for w in q.lower().replace(',',' ').replace('.',' ').replace('(',' ').replace(')',' ').split():
        if len(w) >= 3: word_counter[w] += 1
wb.close()
sections = list(data.keys())
POPULAR_WORDS = sorted(word_counter.keys(), key=lambda w: -word_counter[w])[:500]

class MainApp(App):
    def build(self):
        self.title = 'Bank'
        self.current_section = None
        self.current_page = 0
        self.results = []
        self.mode = 'menu'
        self._t = None
        root = BoxLayout(orientation='vertical', padding=dp(8), spacing=dp(5))
        self.title_label = Label(text='Choose section', font_size=sp(19), size_hint_y=None, height=dp(40), bold=True)
        root.add_widget(self.title_label)
        sr = BoxLayout(size_hint_y=None, height=dp(52), spacing=dp(6))
        self.search_input = TextInput(hint_text='Search...', multiline=False, font_size=sp(17), size_hint_x=0.7)
        self.search_input.bind(text=self.on_search_text)
        bc = Button(text='X', font_size=sp(22), size_hint_x=0.15)
        bc.bind(on_press=self.clear_search)
        bb = Button(text='=', font_size=sp(22), size_hint_x=0.15)
        bb.bind(on_press=self.show_sections)
        sr.add_widget(self.search_input)
        sr.add_widget(bc)
        sr.add_widget(bb)
        root.add_widget(sr)
        self.sugg_box = BoxLayout(orientation='vertical', size_hint_y=None, height=0, spacing=dp(2))
        root.add_widget(self.sugg_box)
        for rw in (QUICK_ROW_1, QUICK_ROW_2, QUICK_ROW_3):
            qr = BoxLayout(size_hint_y=None, height=dp(38), spacing=dp(4))
            for w in rw:
                b = Button(text=w, font_size=sp(13))
                b.bind(on_press=lambda x, ww=w: self.quick_search(ww))
                qr.add_widget(b)
            root.add_widget(qr)
        self.scroll = ScrollView()
        self.content = GridLayout(cols=1, spacing=dp(6), size_hint_y=None, padding=dp(4))
        self.content.bind(minimum_height=self.content.setter('height'))
        self.scroll.add_widget(self.content)
        root.add_widget(self.scroll)
        nav = BoxLayout(size_hint_y=None, height=dp(50), spacing=dp(6))
        self.btn_prev = Button(text='<', font_size=sp(16))
        self.btn_prev.bind(on_press=self.prev_page)
        self.btn_next = Button(text='>', font_size=sp(16))
        self.btn_next.bind(on_press=self.next_page)
        nav.add_widget(self.btn_prev)
        nav.add_widget(self.btn_next)
        root.add_widget(nav)
        self.show_sections()
        return root
    def show_sections(self, *a):
        self.mode='menu'; self.search_input.text=''; self.hide_sugg()
        self.title_label.text='Choose section'; self.content.clear_widgets()
        for i, sec in enumerate(sections, 1):
            b = Button(text='['+str(i)+'] '+sec+' ('+str(len(data[sec]))+')', font_size=sp(18), size_hint_y=None, height=dp(70))
            b.bind(on_press=lambda x, s=sec: self.open_section(s))
            self.content.add_widget(b)
        self.scroll.scroll_y=1; self.set_nav(False, False)
    def open_section(self, s):
        self.mode='section'; self.current_section=s; self.current_page=0
        self.search_input.text=''; self.hide_sugg(); self.render_page()
    def hide_sugg(self):
        self.sugg_box.clear_widgets(); self.sugg_box.height=0
    def show_sugg(self, ms):
        self.sugg_box.clear_widgets()
        if not ms: self.sugg_box.height=0; return
        for w in ms[:MAX_SUGGESTIONS]:
            b = Button(text='  '+w+' ('+str(word_counter.get(w,0))+')', font_size=sp(15), size_hint_y=None, height=dp(34))
            b.bind(on_press=lambda x, ww=w: self.pick_sugg(ww))
            self.sugg_box.add_widget(b)
        n=min(len(ms),MAX_SUGGESTIONS); self.sugg_box.height=n*dp(34)+(n-1)*dp(2)
    def pick_sugg(self, w):
        self.search_input.text=w; self.hide_sugg(); self.run_search(w)
    def on_search_text(self, inst, val):
        if self._t: self._t.cancel()
        t = val.strip().lower()
        if len(t)<2:
            self.hide_sugg()
            if self.mode=='search': self.show_sections()
            return
        ms = [w for w in POPULAR_WORDS if w.startswith(t)]
        self.show_sugg(ms)
        self._t = Clock.schedule_once(lambda dt: self.run_search(t), 0.35)
    def run_search(self, q):
        ql=q.lower().strip(); self.mode='search'; self.current_page=0; self.results=[]
        for sec, c, qq, a in all_items:
            if ql in qq.lower() or ql in a.lower() or ql in c.lower() or ql in sec.lower():
                self.results.append((sec,c,qq,a))
        self.render_search()
    def quick_search(self, w):
        self.search_input.text=w; self.hide_sugg(); self.run_search(w)
    def clear_search(self, *a):
        self.search_input.text=''; self.hide_sugg(); self.show_sections()
    def _item(self, sec, c, q, a, show_sec):
        if show_sec:
            self.content.add_widget(Label(text='Section: '+sec, font_size=sp(13), size_hint_y=None, height=dp(22), halign='left', text_size=(self.scroll.width-dp(30),None), padding=(dp(10),0)))
        self.content.add_widget(Label(text='Code: '+c, font_size=sp(14), size_hint_y=None, height=dp(24), halign='left', text_size=(self.scroll.width-dp(30),None), padding=(dp(10),0)))
        ql = Label(text='Q: '+q, font_size=sp(17), size_hint_y=None, halign='left', valign='top', text_size=(self.scroll.width-dp(30),None), padding=(dp(10),dp(8)))
        ql.bind(texture_size=lambda l, v: setattr(l, 'height', v[1]+dp(20)))
        self.content.add_widget(ql)
        al = Label(text='A: '+a, font_size=sp(17), size_hint_y=None, halign='left', valign='top', text_size=(self.scroll.width-dp(30),None), padding=(dp(10),dp(8)))
        al.bind(texture_size=lambda l, v: setattr(l, 'height', v[1]+dp(20)))
        self.content.add_widget(al)
        self.content.add_widget(Label(text='-'*20, font_size=sp(12), size_hint_y=None, height=dp(20)))
    def render_search(self):
        t=len(self.results); tp=max(1,(t+PAGE_SIZE-1)//PAGE_SIZE)
        self.title_label.text='Found: '+str(t)+' ('+str(self.current_page+1)+'/'+str(tp)+')'
        self.content.clear_widgets()
        if t==0:
            self.content.add_widget(Label(text='Not found', font_size=sp(18), size_hint_y=None, height=dp(80)))
            self.set_nav(False,False); return
        s=self.current_page*PAGE_SIZE; e=min(s+PAGE_SIZE,t)
        for i in range(s,e):
            sec,c,q,a=self.results[i]; self._item(sec,c,q,a,True)
        self.set_nav(self.current_page>0, self.current_page<tp-1); self.scroll.scroll_y=1
    def render_page(self):
        sec=self.current_section; qs=data[sec]; t=len(qs); tp=(t+PAGE_SIZE-1)//PAGE_SIZE
        self.title_label.text=sec+' ('+str(self.current_page+1)+'/'+str(tp)+')'
        self.content.clear_widgets()
        s=self.current_page*PAGE_SIZE; e=min(s+PAGE_SIZE,t)
        for i in range(s,e):
            c,q,a=qs[i]
            self.content.add_widget(Label(text='Q '+str(i+1)+':', font_size=sp(17), size_hint_y=None, height=dp(24), bold=True, halign='left', text_size=(self.scroll.width-dp(30),None), padding=(dp(10),0)))
            self._item('',c,q,a,False)
        self.set_nav(self.current_page>0, self.current_page<tp-1); self.scroll.scroll_y=1
    def next_page(self, *a):
        if self.mode=='section': t=len(data[self.current_section])
        elif self.mode=='search': t=len(self.results)
        else: return
        tp=(t+PAGE_SIZE-1)//PAGE_SIZE
        if self.current_page<tp-1:
            self.current_page+=1
            if self.mode=='section': self.render_page()
            else: self.render_search()
    def prev_page(self, *a):
        if self.current_page>0:
            self.current_page-=1
            if self.mode=='section': self.render_page()
            elif self.mode=='search': self.render_search()
    def set_nav(self, p, n):
        self.btn_prev.disabled=not p; self.btn_prev.opacity=1 if p else 0.3
        self.btn_next.disabled=not n; self.btn_next.opacity=1 if n else 0.3

if __name__ == '__main__':
    MainApp().run()
